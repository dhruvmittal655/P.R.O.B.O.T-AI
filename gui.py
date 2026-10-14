import tkinter as tk
import customtkinter as ctk
import threading
import time
from tkinter import ttk
from tkinter import filedialog
from explorer import ResearchExplorer
from tkinter import messagebox    
from reportlab.platypus import SimpleDocTemplate, Paragraph
from reportlab.lib.styles import getSampleStyleSheet
from tkinter import messagebox
from downloader import Downloader
from database import Database
from source_manager import SourceManager
from research_downloader import ResearchDownloader
from datetime import datetime
from searcher import Searcher
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
from tkinter import simpledialog
import re
from image_loader import ImageLoader
from research_downloader import ResearchDownloader
from datetime import datetime
import os
from semantic_search import SemanticSearch
from notes_generator import NotesGenerator
class SmartSearchGUI:

    def choose_result(self, results):

     window = tk.Toplevel(self.window)

     window.title("Select Article")

     window.geometry("450x350")

     label = tk.Label(window, text="Select the article you want")

     label.pack(pady=10)

     listbox = tk.Listbox(window, width=60, height=12)

     listbox.pack(pady=10)

     for item in results:
        listbox.insert(tk.END, item)

     selected = {"title": None}

     def confirm():

        if listbox.curselection():

            selected["title"] = listbox.get(listbox.curselection())

            window.destroy()

     button = ctk.CTkButton(window, text="Open", command=confirm)

     button.pack(pady=10)

     window.grab_set()

     self.window.wait_window(window)

     return selected["title"]

    def __init__(self):
        ctk.set_appearance_mode("Light")
        ctk.set_default_color_theme("blue")
# ---------------- WINDOW ---------------- #

        self.window = ctk.CTk()
        self.window.title("SmartSearch AI")
        self.window.geometry("1280x900")
        self.window.minsize(1100, 750)

# ---------------- DATABASE ---------------- #

        self.database = Database()

# ---------------- SEARCH ---------------- #

        self.searcher = Searcher()
        self.downloader = Downloader()
        self.image_loader = ImageLoader()

# ---------------- ML ---------------- #

        self.semantic = SemanticSearch()
        self.notes = NotesGenerator()
# ---------------- VARIABLES ---------------- #
        from ai_assistant import SmartAIAssistant
        self.explorer = ResearchExplorer()
        self.ai = SmartAIAssistant(self.database)
        self.ai.current_article = None
        self.research_mode = False
        self.research_topic = ""
        self.dark_mode = False
        self.search_after_id = None

# ---------------- SCROLLABLE MAIN FRAME ---------------- #
        self.current_image = None
        self.main_frame = ctk.CTkScrollableFrame(
    self.window,
    fg_color="transparent"
)
        self.main_frame.pack(
    fill="both",
    expand=True
)
# ---------------- BUILD GUI ---------------- #
        
        self.create_widgets()
        self.create_popup_menu()

        # ---------- Background AI Model Pre-warming ----------
        import threading
        threading.Thread(
            target=lambda: (self.ai.load_nlp(), self.semantic.load_model()),
            daemon=True
        ).start()

        self.window.mainloop()
    def create_widgets(self):
     top_frame = ctk.CTkFrame(
    self.main_frame,
    fg_color="transparent"
)

     top_frame.pack(
    fill="x",
    padx=20,
    pady=(10,0)
)
     self.menu_button = ctk.CTkButton(
    top_frame,
    text="☰",
    width=45,
    height=40,
    font=("Segoe UI",22),
    command=self.show_menu
)

     self.menu_button.pack(
    side="left",
    padx=(0,10)
)
    # ---------------- Window Title ---------------- #
     title = ctk.CTkLabel(
    top_frame,
    text="🔍 P.R.O.B.O.T",
    font=("Segoe UI", 34, "bold"),
    text_color="#053842"
)
     title.pack(
    side="left",
    expand=True,
    padx=10,
    pady=10
)
     ctk.CTkLabel(
    top_frame,
    text="     "
).pack(side="right")

     search_frame = ctk.CTkFrame(
    self.main_frame,
    corner_radius=15,
    fg_color=("white", "#2B2B2B")
)

     search_frame.pack(
    
    pady=15
)

    # ---------------- Search Box ---------------- #
     self.search_box = ctk.CTkEntry(
    search_frame,
    width=600,
    height=45,
    corner_radius=0,
    border_width = 0,
    placeholder_text="Search anything...",
    font=("Segoe UI",16)
)
     self.search_box.pack(
    side="left",

    padx=(0)
)
     self.search_button = ctk.CTkButton(
    search_frame,
    text="🔍 Search",
    command=self.search_article,
    width=160,
    height=45,
    corner_radius=0,
    border_width = 0,
    fg_color="#053842",
hover_color="#5A88A4",
    font=("Segoe UI", 15, "bold")
)

     self.search_button.pack(side="left", padx=(0))
    
     self.search_box.bind(
    "<Return>",
    lambda event: self.search_article()
)
     self.window.bind(
    "<Escape>",
    lambda event: self.hide_suggestions()
)
    # ---------------- Suggestions ---------------- #

     self.suggestion_frame = ctk.CTkFrame(
    self.window,
    width=700,
    height=170,
    corner_radius=12,
    border_width=1,
    fg_color=("white", "#2B2B2B")
)

     self.suggestion_frame.place_forget()

     self.search_box.bind(
    "<KeyRelease>",
    self.schedule_suggestions
)
     # ---------------- Button Bar ---------------- #
     self.button_frame = ctk.CTkFrame(
        self.main_frame,
        fg_color="transparent"
     )
     self.button_frame.pack(
        pady=10
     )

     ctk.CTkButton(
        self.button_frame,
        text="🔬 Research",
        command=self.build_research_pack,
        width=120,
        fg_color="#053842",
        hover_color="#5A88A4"
     ).pack(side="left", padx=5)

     ctk.CTkButton(
        self.button_frame,
        text="↩ Normal Search",
        command=self.exit_research_mode,
        width=145,
        fg_color="#053842",
        hover_color="#5A88A4"
     ).pack(side="left", padx=5)

     ctk.CTkButton(
        self.button_frame,
        text="💬 AI Chat",
        command=self.open_ai_chat,
        width=110,
        fg_color="#053842",
        hover_color="#5A88A4"
     ).pack(side="left", padx=5)

     ctk.CTkButton(
        self.button_frame,
        text="🧠 Study Notes",
        command=self.generate_notes,
        width=120,
        fg_color="#053842",
        hover_color="#5A88A4"
     ).pack(side="left", padx=5)

     ctk.CTkButton(
        self.button_frame,
        text="📄 Upload PDF",
        command=self.upload_external_pdf,
        width=120,
        fg_color="#1E5631",
        hover_color="#2D8249"
     ).pack(side="left", padx=5)

     ctk.CTkButton(
        self.button_frame,
        text="📂 Saved PDF Library",
        command=self.show_downloaded_pdfs_library,
        width=140,
        fg_color="#005B96",
        hover_color="#003B63"
     ).pack(side="left", padx=5)

     ctk.CTkButton(
        self.button_frame,
        text="🎓 Cite Document",
        command=self.show_citation_dialog,
        width=135,
        fg_color="#053842",
        hover_color="#5A88A4"
     ).pack(side="left", padx=5)

     self.workspace_frame = ctk.CTkFrame(
        self.main_frame,
        fg_color="transparent"
     )

     self.workspace_frame.pack(
        fill="both",
        expand=True,
        padx=20,
        pady=10
     )


     # ---------- Results Container ----------

     self.results_container = ctk.CTkFrame(
    self.workspace_frame,
    fg_color="transparent"
)   
     self.results_container.pack(
     side = "right",
     fill="both",
     expand=True)

     self.explorer_frame = ctk.CTkFrame(
    self.workspace_frame,
    width=280,
    fg_color=("white", "#2B2B2B"),
    corner_radius=15,
    border_width=1
)
     self.explorer_frame.pack(
        side="left",
        fill="y",
        padx=(0, 5),
        before=self.results_container
     )

     self.explorer_label = ctk.CTkLabel(
         self.explorer_frame,
         text="📖 Research Explorer",
         font=("Segoe UI", 18, "bold")
     )
     
     self.explorer_label.pack(
         pady=10
     )
     
     self.explorer_buttons = ctk.CTkScrollableFrame(
         self.explorer_frame,
         width=200,
         height=500
     )
     
     self.explorer_buttons.pack(
         fill="both",
         expand=True,
         padx=5,
         pady=5
     )

     self.delete_pack_btn = ctk.CTkButton(
        self.explorer_frame,
        text="🗑️ Delete Research Pack",
        fg_color="#8B0000",
        hover_color="#A52A2A",
        height=32,
        font=("Segoe UI", 12, "bold"),
        command=self.delete_current_research_pack
     )
     self.delete_pack_btn.pack(fill="x", padx=10, pady=(5, 10))

     self.result_frame = ctk.CTkFrame(
    self.results_container,
    corner_radius=18,
    border_width=1
)
     
     self.result_frame.pack(
    side="left",
    fill="both",
    expand=True,
    padx=5,
    pady=5
)

     
# ---------------- Image ---------------- #
     self.image_frame = ctk.CTkFrame(
    self.result_frame,
    corner_radius=15
)

     self.image_frame.pack(
    pady=(0,10),
    anchor="center"
)

     self.image_label = ctk.CTkLabel(
    self.image_frame,
    text=""
)

     self.image_label.pack(
    padx=7,
    pady=7
)
#------------result frame------------#
     
     

# ---------------- Research Explorer ---------------- #

    

# ---------------- Related Articles ---------------- #

     self.related_frame = ctk.CTkFrame(
    self.main_frame,
    corner_radius=10,
   
)

     self.related_frame.pack(
    fill="x",
    padx=20,
    pady=(10,10)
)

     self.related_frame.configure(height=120)
     self.related_frame.pack_propagate(False)
     

     self.related_title = ctk.CTkLabel(
    self.related_frame,
    text="🧠 Related Articles",
    font=("Segoe UI",16,"bold")
)

     self.related_title.pack(
    anchor="w",
    padx=10,
    pady=(8,5)
)

     self.related_list = ctk.CTkFrame(
    self.related_frame,
    fg_color="transparent"
)

     self.related_list.pack(
    fill="x",
    padx=10,
    pady=(0,10)
)
# ---------------- Status Bar ---------------- #
     self.connection_label = ctk.CTkLabel(
    self.main_frame,
    text="🟢 Online",
    font=("Segoe UI", 12, "bold")
)

     self.connection_label.pack(
    pady=(0, 5)
)
     self.status = ctk.CTkLabel(
    self.main_frame,
    text="Ready",
    anchor="w",
    height=28,
    font=("Segoe UI", 12),
    fg_color="#F2F5F2",
    text_color="black",
    corner_radius=0
)
     self.status.pack(
    fill="x",
    pady=(10, 5)
)


    
     scrollbar = tk.Scrollbar(self.result_frame)

     self.result_box = tk.Text(
    self.result_frame,
    wrap="word",
    font=("Segoe UI", 15),
    bd=0,
    relief="flat",
    padx=30,
    pady=30,
    yscrollcommand=scrollbar.set
)

     scrollbar.config(command=self.result_box.yview)

     scrollbar.pack(
    side="right",
    fill="y"
)

     self.result_box.pack(
    side="left",
    fill="both",
    expand=True
)

     def open_clicked_link(event):
         try:
             import webbrowser, re
             index = self.result_box.index(f"@{event.x},{event.y}")
             line_text = self.result_box.get(f"{index} linestart", f"{index} lineend")
             urls = re.findall(r'https?://[^\s]+|file:///[^\s]+', line_text)
             if urls:
                 target_url = urls[0].rstrip('.)],;')
                 if "Download PDF" in line_text:
                     self.download_and_open_pdf(target_url, "Research_Paper")
                 else:
                     webbrowser.open(target_url)
                     self.status.configure(text=f"Opening link: {target_url[:50]}...")
         except Exception:
             pass

     self.result_box.bind("<Button-1>", open_clicked_link)
     self.result_box.bind("<Double-Button-1>", open_clicked_link)






     self.window.bind("<Button-1>", self.hide_suggestions)

    
     
    def search_article(self):
     self.prepare_result_box()
     self.research_mode = False
     self.research_topic = None

     self.status.configure(
    text="Normal Mode"
)
     self.window.configure(cursor="watch")
     self.window.update()
     self.search_button.configure(state="disabled")
     self.status.configure(text="Searching...")
     self.window.update_idletasks()
     online = self.check_connection()
     raw_query = self.search_box.get().strip()
     if raw_query == "":
      self.status.configure(text="Enter a topic or question.")
      self.window.configure(cursor="")
      self.search_button.configure(state="normal")
      return

      # ---------- Smart Intent & DOI Query Resolution ----------
     if "10." in raw_query and "/" in raw_query:
         self.status.configure(text=f"Ingesting paper via DOI '{raw_query}'...")
         self.window.update()
         from doi_fetcher import DOIFetcher
         fetcher = DOIFetcher()
         doi_data = fetcher.fetch_by_doi(raw_query)
         if doi_data:
             title = doi_data["title"]
             topic = f"[DOI] {title}"
             now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
             self.database.save_or_update_article(topic, doi_data["abstract"])
             self.database.save_research_source(
                 topic,
                 "DOI Paper Ingest",
                 doi_data["doi_url"],
                 doi_data["content"],
                 now_str,
                 doi_data["authors"],
                 doi_data["pub_date"],
                 ""
             )
             self.research_mode = True
             self.research_topic = topic
             self.ai.current_article = topic
             self.update_reader_text(doi_data["content"])
             sources = [("DOI Paper Ingest", doi_data["doi_url"], doi_data["content"])]
             self.show_research_sources(sources)
             self.status.configure(text=f"🟢 Ingested DOI Paper: {title}")
             self.window.configure(cursor="")
             self.search_button.configure(state="normal")
             return

     topic = raw_query
     word_count = len(raw_query.split())
     is_question_format = "?" in raw_query or any(raw_query.lower().startswith(q) for q in ["what", "who", "when", "where", "why", "how", "explain", "tell me", "can you", "is", "are", "describe"])

     if word_count > 3 or is_question_format:
        all_titles = self.database.get_all_titles()
        best_db = self.semantic.best_matching_title(all_titles, raw_query) if all_titles else None
        if best_db and self.database.article_exists(best_db):
            topic = best_db
        else:
            extracted = self.ai.extract_topic(raw_query)
            if extracted and len(extracted.strip()) > 1:
                topic = extracted.strip()

     if not online:
       self.status.configure(text="Offline Mode")
       local_titles = [row[0] for row in self.database.search_titles(topic)]
       exact_title = next(
           (title for title in local_titles if title.lower() == topic.lower()),
           None
       )

       if exact_title:
        selected_title = exact_title
       elif local_titles:
        selected_title = self.choose_result(local_titles)
       else:
        selected_title = None

       if not selected_title:
        self.result_box.delete("1.0", "end")
        self.result_box.insert(
            "end",
            "No internet connection and no matching saved article was found."
        )
        self.search_button.configure(state="normal")
        self.window.configure(cursor="")
        return

       article = self.database.get_article(selected_title)

       if not article:
        self.result_box.delete("1.0", "end")
        self.result_box.insert("end", "Saved article could not be loaded.")
        self.search_button.configure(state="normal")
        self.window.configure(cursor="")
        return

       self.ai.set_article(selected_title, article[2])
       self.database.update_views(selected_title)
       self.database.update_last_opened(selected_title)
       self.result_box.delete("1.0", "end")
       self.result_box.insert("end", article[2])
       self.show_image(self.database.get_image_path(selected_title))
       self.show_related_articles([])
       self.status.configure(text=f"Loaded offline: {selected_title}")
       self.search_button.configure(state="normal")
       self.window.configure(cursor="")
       return

     
     self.suggestion_frame.place_forget()

    

     self.result_box.delete("1.0", tk.END)

    # ---------------- Semantic Suggestions ---------------- #

     titles = self.database.get_all_titles()

     matches = self.semantic.top_matches(
        topic,
        titles
    )

     

     

     

     if self.database.article_exists(topic):
        article = self.database.get_article(topic)
        if article and article[2] and not article[2].startswith("Research Pack for"):
            self.ai.set_article(topic, article[2])
            self.ai.current_article = topic
            self.research_mode = False
            self.database.update_views(topic)
            self.database.update_last_opened(topic)

            self.result_box.insert(tk.END, article[2])

            sections = self.explorer.extract_sections(article[2])
            if len(sections) > 1 or (len(sections) == 1 and sections[0]["title"] != "Full Article"):
                self.show_research_explorer(sections)
            else:
                self.explorer_frame.pack_forget()

            image_path = self.database.get_image_path(topic)
            if not image_path or not os.path.exists(image_path):
                image_url = self.image_loader.get_image_url(topic)
                if image_url:
                    image_path = self.image_loader.download_image(image_url, topic)

            if image_path:
                self.database.update_image_path(topic, image_path)

            self.show_image(image_path)
            history_titles = [row[0] for row in self.database.get_history()]
            related = self.semantic.top_matches(topic, history_titles)
            self.show_related_articles(related)

            self.status.configure(text="Loaded from database.")
            self.window.configure(cursor="")
            self.search_button.configure(state="normal")
            return

    # ---------------- WIKIPEDIA SEARCH ---------------- #

     results = self.searcher.search(topic)

     if not results:

        self.result_box.insert(
            tk.END,
            "No matching article found."
        )

        self.status.configure(
            text="Search failed."
        )
        self.window.configure(cursor="")
        self.search_button.configure(state="normal")
        return

     best_title = self.choose_result(results)

     if best_title is None:
        self.window.configure(cursor="")
        self.search_button.configure(state="normal")
        return

     summary = self.downloader.get_summary(best_title)

     if not summary:

        self.result_box.insert(
            tk.END,
            "Unable to download article."
        )

        self.status.configure(
            text="Download failed."
        )
        self.window.configure(cursor="")
        self.search_button.configure(state="normal")
        return

     image_path = None

     image_url = self.image_loader.get_image_url(best_title)

     if image_url:

      image_path = self.image_loader.download_image(
        image_url,
        best_title
    )

# Show image whether it exists or not.
# show_image() will handle "No Image" itself.
     self.show_image(image_path)

     from datetime import datetime

     current = datetime.now().strftime(
        "%Y-%m-%d %H:%M:%S"
    )

     self.database.save_or_update_article(
        best_title,
        summary,
        image_path,
        current,
        current
     )
     self.ai.set_article(
    best_title,
    summary
)
     print("AI topic after search:", self.ai.current_topic)
     self.ai.current_article = best_title
     self.research_mode = False
     self.result_box.insert(
        tk.END,
        summary
    )
     sections = self.explorer.extract_sections(summary)
     if len(sections) > 1 or (len(sections) == 1 and sections[0]["title"] != "Full Article"):
         self.show_research_explorer(sections)
     else:
         self.explorer_frame.pack_forget()

     print("\n========== Sections Found ==========\n")

     for section in sections:
      print(section["title"])
     history_titles = [
        row[0]
        for row in self.database.get_history()
    ]

     related = self.semantic.top_matches(
        best_title,
        history_titles
    )

     self.show_related_articles(related)

     self.status.configure(
        text=f"Downloaded: {best_title}"
    )
     self.window.configure(cursor="")
     self.search_button.configure(state="normal")
      
    def load_topic_from_db(self, title):
        title = str(title).strip()
        if not title:
            return

        self.prepare_result_box()
        self.search_box.delete(0, "end")
        self.search_box.insert(0, title)

        self.database.update_views(title)
        self.database.update_last_opened(title)

        if self.database.research_pack_exists(title):
            self.research_mode = True
            self.research_topic = title
            self.ai.current_article = title
            self.show_research_sources(self.get_research_sources_for_explorer(title))
            full_text = self.database.get_full_research(title)
            self.result_box.delete("1.0", tk.END)
            self.result_box.insert(tk.END, full_text)
        else:
            article = self.database.get_article(title)
            if article:
                self.research_mode = False
                self.ai.set_article(title, article[2])
                self.ai.current_article = title
                self.result_box.delete("1.0", tk.END)
                self.result_box.insert(tk.END, article[2])

                sections = self.explorer.extract_sections(article[2])
                if len(sections) > 1 or (len(sections) == 1 and sections[0]["title"] != "Full Article"):
                    self.show_research_explorer(sections)
                else:
                    self.explorer_frame.pack_forget()

        image_path = self.database.get_image_path(title)
        self.show_image(image_path)

        history_titles = [row[0] for row in self.database.get_history()]
        related = self.semantic.top_matches(title, history_titles)
        self.show_related_articles(related)
        self.status.configure(text=f"Loaded from history: {title}")

    def show_history(self):

     from tkinter import ttk, messagebox

     window = tk.Toplevel(self.window)
     window.title("History Manager")
     window.geometry("850x500")

    # =======================
    # TABLE FRAME
    # =======================

     table_frame = tk.Frame(window)
     table_frame.pack(fill="both", expand=True, padx=10, pady=10)

     columns = ("Title", "Last Opened", "Views")

     tree = ttk.Treeview(
        table_frame,
        columns=columns,
        show="headings"
    )

     tree.heading("Title", text="Title")
     tree.heading("Last Opened", text="Last Opened")
     tree.heading("Views", text="Views")

     tree.column("Title", width=420)
     tree.column("Last Opened", width=220)
     tree.column("Views", width=80, anchor="center")

     scrollbar = ttk.Scrollbar(
        table_frame,
        orient="vertical",
        command=tree.yview
    )

     tree.configure(yscrollcommand=scrollbar.set)

     tree.pack(side="left", fill="both", expand=True)

     scrollbar.pack(side="right", fill="y")

    # =======================
    # FUNCTIONS
    # =======================

     def refresh():

        tree.delete(*tree.get_children())

        history = self.database.get_history()

        for row in history:
            tree.insert("", "end", values=row)

     def open_article(event=None):
      selected = tree.selection()
      if not selected:
        return
      values = tree.item(selected[0])["values"]
      title = values[0]
      self.load_topic_from_db(title)
      window.destroy()

     def delete_selected():

        selected = tree.selection()

        if not selected:
            messagebox.showwarning(
                "Delete",
                "Select an article first."
            )
            return

        title = tree.item(selected[0])["values"][0]

        if messagebox.askyesno(
            "Delete",
            f"Delete '{title}'?"
        ):

            self.database.delete_article(title)

            refresh()

     def clear_all():

        if messagebox.askyesno(
            "Clear History",
            "Delete ALL history?"
        ):

            self.database.clear_history()

            refresh()

    # =======================
    # BUTTON FRAME
    # =======================

     button_frame = tk.Frame(window)
     button_frame.pack(fill="x", pady=10)

     tk.Button(
        button_frame,
        text="Open",
        width=15,
        command=open_article
    ).pack(side="left", padx=8)

     tk.Button(
        button_frame,
        text="Delete Selected",
        width=18,
        command=delete_selected
    ).pack(side="left", padx=8)

     tk.Button(
        button_frame,
        text="Clear All",
        width=15,
        command=clear_all
    ).pack(side="left", padx=8)

     tk.Button(
        button_frame,
        text="Refresh",
        width=15,
        command=refresh
    ).pack(side="left", padx=8)

     tk.Button(
        button_frame,
        text="Close",
        width=15,
        command=window.destroy
    ).pack(side="right", padx=8)

     tree.bind("<Double-1>", open_article)

     refresh()
 
 
     

    def add_to_favorites(self):

      if self.ai.current_article is None:

        self.status.configure(text="No article selected.")

        return

      self.database.add_favorite(self.ai.current_article)

      self.status.configure(text="Added to favorites.")
    
    def show_favorites(self):

     favorites = self.database.get_favorites()

     window = tk.Toplevel(self.window)

     window.title("Favorites")

     window.geometry("500x400")

     listbox = tk.Listbox(
        window,
        width=60,
        height=20
    )

     listbox.pack(padx=10, pady=10, fill="both", expand=True)

     if len(favorites) == 0:

        listbox.insert(tk.END, "No favorites added.")

        return

     for article in favorites:

        listbox.insert(tk.END, article[0])

     def open_article(event):

      selection = listbox.curselection()

      if not selection:
        return

      title = listbox.get(selection)

      article = self.database.get_article(title)

      if article:

        # ---------- Text ----------

        self.prepare_result_box()
        self.result_box.delete("1.0", tk.END)
        self.result_box.insert(tk.END, article[2])

        self.ai.current_article = title

        self.database.update_views(title)
        self.database.update_last_opened(title)

      self.load_topic_from_db(title)
      window.destroy()

     listbox.bind("<Double-Button-1>", open_article)
    def show_most_viewed(self):

     articles = self.database.get_most_viewed()

     window = tk.Toplevel(self.window)

     window.title("Most Viewed")

     window.geometry("500x400")

     listbox = tk.Listbox(
        window,
        width=60,
        height=20
     )

     listbox.pack(padx=10, pady=10, fill="both", expand=True)

     if len(articles) == 0:

        listbox.insert(tk.END, "No articles available.")

        return

     for article in articles:

        listbox.insert(
            tk.END,
            f"{article[0]} ({article[1]} views)"
        )
     def open_article(event):

      selection = listbox.curselection()

      if not selection:
        return

      text = listbox.get(selection)
 
      title = text.rsplit(" (", 1)[0]
      self.load_topic_from_db(title)
      window.destroy()

     listbox.bind("<Double-Button-1>", open_article)
    def show_suggestions(self, event=None):

     keyword = self.search_box.get().strip()

     if keyword == "":
        self.suggestion_frame.place_forget()
        return

    # Clear old suggestions
     for widget in self.suggestion_frame.winfo_children():
        widget.destroy()

     results = self.database.search_titles(keyword)

     if not results:
        self.suggestion_frame.place_forget()
        return

     added = set()

     for result in results:

        title = result[0]

        if title in added:
            continue

        added.add(title)

        btn = ctk.CTkButton(
            self.suggestion_frame,
            text="⭐ " + title,
            anchor="w",
            fg_color="transparent",
            hover_color=("gray85", "gray25"),
            text_color=("black", "white"),
            height=28,
            command=lambda t=title: self.select_suggestion(t)
        )

        btn.pack(fill="x", padx=5, pady=2)

     self.window.update_idletasks()

     x = (
        self.search_box.winfo_rootx()
        - self.window.winfo_rootx()
    )

     y = (
        self.search_box.winfo_rooty()
        - self.window.winfo_rooty()
        + self.search_box.winfo_height()
        + 5
    )

     self.suggestion_frame.configure(
        width=self.search_box.winfo_width()
    )

     self.suggestion_frame.place(
        x=x,
        y=y
    )

     self.suggestion_frame.lift()
    def hide_suggestions(self, event=None):

     if event is None:
        self.suggestion_frame.place_forget()
        return

     widget = event.widget

    # Keep suggestions open while typing
     if widget == self.search_box:
        return

    # Keep them open when clicking a suggestion button
     if str(widget).startswith(str(self.suggestion_frame)):
        return

    # Otherwise hide
     self.suggestion_frame.place_forget()
    def select_suggestion(self, title):

     self.search_box.delete(0, "end")
     self.search_box.insert(0, title)

     self.suggestion_frame.place_forget()

     self.search_article()
    def schedule_suggestions(self, event=None):

     if hasattr(self, "search_after_id") and self.search_after_id:
        self.window.after_cancel(self.search_after_id)

     self.search_after_id = self.window.after(
        250,
        self.show_suggestions
    )

    def export_docx(self):
        text = self.result_box.get("1.0", tk.END).strip()
        if text == "":
            self.status.configure(text="Nothing to export.")
            return

        filename = filedialog.asksaveasfilename(
            defaultextension=".docx",
            filetypes=[("Word Document", "*.docx"), ("All Files", "*.*")]
        )
        if filename == "":
            return

        try:
            from docx import Document
            doc = Document()
            doc.add_heading("ProBot AI Research Result", level=1)

            for line in text.split("\n"):
                doc.add_paragraph(line)

            topic = getattr(self.ai, "current_article", None) or getattr(self, "research_topic", None)
            if topic:
                refs = self.database.get_references_for_topic(topic)
                if refs:
                    doc.add_heading("References & Citation Bibliography", level=2)
                    for name, url, downloaded, authors, pub_date in refs:
                        cit = f"• {name} ({pub_date or 'n.d.'}). {authors or 'Author'}. Available at: {url} [Accessed {downloaded or 'Recently'}]."
                        doc.add_paragraph(cit)

            doc.save(filename)
            self.status.configure(text="Word document exported successfully with references.")
        except Exception as e:
            self.status.configure(text=f"Error: {e}")

    def export_pdf(self):
        text = self.result_box.get("1.0", tk.END).strip()
        if text == "":
            self.status.configure(text="Nothing to export.")
            return

        filename = filedialog.asksaveasfilename(
            defaultextension=".pdf",
            filetypes=[("PDF File", "*.pdf")]
        )
        if filename == "":
            return

        try:
            doc = SimpleDocTemplate(filename)
            styles = getSampleStyleSheet()
            story = []
            story.append(Paragraph("<b>ProBot AI Research Result</b>", styles["Heading1"]))

            for line in text.split("\n"):
                story.append(Paragraph(line, styles["BodyText"]))

            topic = getattr(self.ai, "current_article", None) or getattr(self, "research_topic", None)
            if topic:
                refs = self.database.get_references_for_topic(topic)
                if refs:
                    story.append(Paragraph("<b>References & Citation Bibliography</b>", styles["Heading2"]))
                    for name, url, downloaded, authors, pub_date in refs:
                        cit = f"• {name} ({pub_date or 'n.d.'}). {authors or 'Author'}. URL: {url}"
                        story.append(Paragraph(cit, styles["BodyText"]))

            doc.build(story)
            self.status.configure(text="PDF exported successfully with references.")
        except Exception as e:
            self.status.configure(text=f"Error: {e}")

    def export_txt(self):
        text = self.result_box.get("1.0", tk.END).strip()
        if text == "":
            self.status.configure(text="Nothing to export.")
            return

        filename = filedialog.asksaveasfilename(
            defaultextension=".txt",
            filetypes=[("Text File", "*.txt")]
        )
        if filename == "":
            return

        try:
            topic = getattr(self.ai, "current_article", None) or getattr(self, "research_topic", None)
            if topic:
                refs = self.database.get_references_for_topic(topic)
                if refs:
                    text += "\n\n=====================================\nREFERENCES & CITATION BIBLIOGRAPHY\n=====================================\n"
                    for name, url, downloaded, authors, pub_date in refs:
                        text += f"• {name} ({pub_date or 'n.d.'}). {authors or 'Author'}. URL: {url} [Accessed {downloaded or 'Recently'}]\n"

            with open(filename, "w", encoding="utf-8") as file:
                file.write(text)

            self.status.configure(text="TXT exported successfully with references.")
        except Exception as e:
            self.status.configure(text=f"Error: {e}")

    

    def toggle_dark_mode(self):

     if ctk.get_appearance_mode() == "Light":

        ctk.set_appearance_mode("Dark")
        self.status.configure(text="Dark Mode Enabled")

     else:

        ctk.set_appearance_mode("Light")
        self.status.configure(text="Light Mode Enabled")

    

    def show_statistics(self):

     articles = self.database.total_articles()
     favorites = self.database.total_favorites()
     views = self.database.total_views()

     messagebox.showinfo(
        "SmartSearch Statistics",
        f"""
     Articles Saved : {articles}

     Favorites : {favorites}

     Total Views : {views}
     """
    )
    def ask_ai(self):
     print("================================")
     print("Research Mode =", self.research_mode)

     print("Current Topic =", self.ai.current_article)
     print("================================")

     if self.research_mode:

        rows = self.database.get_chunks(
            self.research_topic
        )

        article = []

        for chunk, source in rows:
            article.append(chunk)

     else:

        article = self.ai.current_article

     if not article:
        self.status.configure(
            text="Search an article first."
        )
        return

     question = simpledialog.askstring(
        "Ask AI",
        "Ask a question about this article:"
    )
 
     if not question:
        return

     self.status.configure(
        text="AI is thinking..."
    )

     print("Research Mode:", self.research_mode)

     if self.research_mode:
        print("Chunks:", len(article))

     answer = self.ai.ask(
        article,
        question
    )
     self.prepare_result_box()
     self.result_box.insert(
        tk.END,
        "\n\n=====================================\n"
    )

     self.result_box.insert(
        tk.END,
        "AI Answer:\n\n"
    )

     self.result_box.insert(
        tk.END,
        answer
    )

     self.status.configure(
        text="Ready"
    )    
   
    def generate_notes(self):
     self.prepare_result_box()

     article = self.ai.current_article

     if not article:
        self.status.configure(text="Search an article first.")
        return

     title = self.ai.current_article

     notes = self.notes.generate_notes(title, article)

     self.result_box.delete("1.0", tk.END)

     self.result_box.insert(
        tk.END,
        "==============================\n"
    )
     self.result_box.insert(
        tk.END,
        "      SMART STUDY NOTES\n"
    )
     self.result_box.insert(
        tk.END,
        "==============================\n\n"
    )

    # ---------------- Topic ----------------

     self.result_box.insert(
        tk.END,
        f"TOPIC\n{notes['title']}\n\n"
    )

    # ---------------- Overview ----------------

     self.result_box.insert(
        tk.END,
        "OVERVIEW\n"
    )

     self.result_box.insert(
        tk.END,
        notes["overview"] + "\n\n"
    )

    # ---------------- Key Points ----------------

     self.result_box.insert(
        tk.END,
        "KEY POINTS\n"
    )

     for point in notes["key_points"]:
        self.result_box.insert(
            tk.END,
            f"• {point}\n"
        )

    # ---------------- Important Facts ----------------

     self.result_box.insert(
        tk.END,
        "\nIMPORTANT FACTS\n"
    )

     for fact in notes["facts"]:
        self.result_box.insert(
            tk.END,
            f"• {fact}\n"
        )
     self.result_box.insert(
    tk.END,
    "\nCONCEPT MAP\n\n"
)

     self.result_box.insert(
    tk.END,
    notes["concept_map"]
)

     self.result_box.insert(
    tk.END,
    "\n"
)
    # ---------------- Did You Know ----------------

     self.result_box.insert(
    tk.END,
    "\nDID YOU KNOW?\n\n"
)

     for fact in notes["did_you_know"]:
      self.result_box.insert(
        tk.END,
        f"• {fact}\n"
    )
     

    # ---------------- Concept Map ----------------

     self.result_box.insert(
    tk.END,
    "\n" + notes["concept_map"] + "\n"
)
    # ---------------- Exam Questions ----------------

     self.result_box.insert(
        tk.END,
        "\nPOSSIBLE EXAM QUESTIONS\n"
    )

     for i, question in enumerate(notes["questions"], 1):
        self.result_box.insert(
            tk.END,
            f"{i}. {question}\n"
        )

    # ---------------- Quick Revision ----------------

     self.result_box.insert(
        tk.END,
        "\nQUICK REVISION\n\n"
    )

     self.result_box.insert(
        tk.END,
        notes["revision"] + "\n"
    )

    # ---------------- Reading Stats ----------------

     self.result_box.insert(
        tk.END,
        "\n--------------------------------\n"
    )

     self.result_box.insert(
        tk.END,
        f"Estimated Reading Time : {notes['reading_time']} minute(s)\n"
    )

     self.result_box.insert(
        tk.END,
        f"Word Count : {notes['word_count']}\n"
    )

     self.status.configure(
        text="Study notes generated."
    )
    def show_image(self, image_path):

     from PIL import Image
     import customtkinter as ctk
     import os

    # Clear previous image
     self.image_label.destroy()

     self.image_label = ctk.CTkLabel(
    self.image_frame,
    text=""
)

     self.image_label.pack(padx=7, pady=7)
     self.current_image = None

     if not image_path or not os.path.exists(image_path):
        self.image_label.configure(text="No Image")
        return

     try:

        image = Image.open(image_path)
        image.thumbnail((140, 160))

        self.current_image = ctk.CTkImage(
            light_image=image,
            dark_image=image,
            size=image.size
        )

        

        self.image_label.configure(
            image=self.current_image,
            text=""
        )

     except Exception as e:

        print("Show Image Error:", e)

        self.current_image = None

        self.image_label.configure(
            image=None,
            text="No Image"
        )
   
    def show_related_articles(self, related_articles):

    # Remove old buttons
     for widget in self.related_list.winfo_children():
        widget.destroy()

    # Create new buttons
     for article, score in related_articles:

        btn = ctk.CTkButton(
            self.related_list,
            text=article,
            width=150,
            fg_color="#053842",
            hover_color="#5A88A4",
            command=lambda a=article: self.load_related_article(a)
        )

        btn.pack(side="left", padx=5, pady=5)

    def load_related_article(self, article):

     self.search_box.delete(0, "end")
     self.search_box.insert(0, article)

     self.search_article()
    
    def check_connection(self):

     import requests

     try:

        requests.get(
            "https://www.google.com",
            timeout=3
        )

        self.connection_label.configure(
            text="🟢 Online",
        )

        return True

     except:

        self.connection_label.configure(
            text="🔴 Offline"
        )

        return False

    def split_article(self, article):

     article = article.replace("\n", " ")

     sentences = re.split(
        r'(?<=[.!?])\s+',
        article
    )

     chunks = []

     current = ""

     for sentence in sentences:

        if len(current) + len(sentence) < 700:

            current += " " + sentence

        else:

            chunks.append(current.strip())

            current = sentence

     if current:
        chunks.append(current.strip())

     return chunks

    def detect_question_type(self, question):

     question = question.lower()

     if "who" in question:
        return "who"

     if "when" in question:
        return "when"

     if "where" in question:
        return "where"

     if "why" in question:
        return "why"

     if "how" in question:
        return "how"

     if "what" in question:
        return "what"

     return "general"

    def prepare_context(self, article):

     if article is None:
        return ""

     article = article.replace("\n", " ")
 
     sentences = article.split(". ")

     context = ""

     for sentence in sentences[:40]:
        context += sentence.strip() + ". "

     return context

    def open_ai_chat(self):

     import customtkinter as ctk

    def open_ai_chat(self):
     import customtkinter as ctk
     import threading
     import os

     win = ctk.CTkToplevel(self.window)
     win.title("🤖 SmartSearch AI Assistant")
     win.geometry("820x680")
     win.minsize(700, 550)

     # ---------------- Top Header ---------------- #
     header_frame = ctk.CTkFrame(win, fg_color=("#E5E7EB", "#1F2937"), corner_radius=10)
     header_frame.pack(fill="x", padx=15, pady=(15, 5))

     header_title = ctk.CTkLabel(
        header_frame,
        text="🤖 SmartSearch AI Assistant",
        font=("Segoe UI", 15, "bold"),
        text_color=("#053842", "#E5E7EB")
     )
     header_title.pack(side="left", padx=12, pady=10)

     def clear_chat():
        for widget in chat_scroll.winfo_children():
            widget.destroy()
        self.ai.chat_history.clear()

     def export_chat():
        content = []
        for role, msg in self.ai.chat_history:
            content.append(f"### {role}\n{msg}\n")
        if not content:
            messagebox.showinfo("Export Chat", "No chat history to export.")
            return

        filename = filedialog.asksaveasfilename(
            defaultextension=".md",
            filetypes=[("Markdown File", "*.md"), ("Text File", "*.txt")]
        )
        if filename:
            try:
                with open(filename, "w", encoding="utf-8") as f:
                    f.write(f"# SmartSearch AI Chat History\nScope: {self.chat_scope_mode}\n\n" + "\n".join(content))
                messagebox.showinfo("Export Chat", "Chat history exported successfully!")
            except Exception as ex:
                messagebox.showerror("Export Chat Error", str(ex))

     ctk.CTkButton(
        header_frame,
        text="🗑️ Clear",
        width=85,
        fg_color="#8B0000",
        hover_color="#A52A2A",
        command=clear_chat
     ).pack(side="right", padx=(5, 10), pady=8)

     ctk.CTkButton(
        header_frame,
        text="💾 Export",
        width=85,
        fg_color="#053842",
        hover_color="#5A88A4",
        command=export_chat
     ).pack(side="right", padx=3, pady=8)

     # ---------------- Quick Prompt Chips ---------------- #
     chips_frame = ctk.CTkFrame(win, fg_color="transparent")
     chips_frame.pack(fill="x", padx=15, pady=5)

     def send_prompt(prompt_text):
        question_box.delete(0, "end")
        question_box.insert(0, prompt_text)
        send()

     prompts = [
        ("💡 3-Point Summary", "Summarize the key information in 3 bullet points."),
        ("❓ 3 Exam Questions", "Generate 3 key exam questions."),
        ("🔍 Compare Topics", "Compare the main topics stored in the library."),
        ("⚡ Key Takeaways", "What are the main key takeaways?")
     ]

     for label, p_text in prompts:
        ctk.CTkButton(
            chips_frame,
            text=label,
            height=26,
            fg_color=("#D1D5DB", "#374151"),
            hover_color=("#9CA3AF", "#4B5563"),
            text_color=("black", "white"),
            font=("Segoe UI", 12),
            command=lambda t=p_text: send_prompt(t)
        ).pack(side="left", padx=4)

     # ---------------- Chat Scope Selector ---------------- #
     scope_frame = ctk.CTkFrame(win, fg_color="transparent")
     scope_frame.pack(fill="x", padx=15, pady=(2, 6))

     ctk.CTkLabel(scope_frame, text="Chat Scope:", font=("Segoe UI", 12, "bold")).pack(side="left", padx=(0, 8))

     def set_scope(val):
         self.chat_scope_mode = "current" if "Current" in val else "global"
         self.status.configure(text=f"AI Chat Scope: {val}")

     scope_seg = ctk.CTkSegmentedButton(
         scope_frame,
         values=["🎯 Current File Only", "🌐 Entire Research Library"],
         command=set_scope
     )
     scope_seg.set("🎯 Current File Only" if getattr(self, "chat_scope_mode", "current") == "current" else "🌐 Entire Research Library")
     scope_seg.pack(side="left", fill="x", expand=True)

     # ---------------- Scrollable Chat Box ---------------- #
     chat_scroll = ctk.CTkScrollableFrame(win, fg_color="transparent")
     chat_scroll.pack(fill="both", expand=True, padx=15, pady=5)

     # ---------------- Input Frame ---------------- #
     input_frame = ctk.CTkFrame(win, fg_color="transparent")
     input_frame.pack(fill="x", padx=15, pady=(5, 15))

     question_box = ctk.CTkEntry(
        input_frame,
        placeholder_text="Ask anything about the currently open file or library...",
        font=("Segoe UI", 14),
        height=40
     )
     question_box.pack(side="left", fill="x", expand=True, padx=(0, 10))

     def copy_text(text_val, btn):
        self.window.clipboard_clear()
        self.window.clipboard_append(text_val)
        orig = btn.cget("text")
        btn.configure(text="✓ Copied!")
        win.after(1500, lambda: btn.configure(text=orig))

     def speak_text(text_val):
        def run_speech():
            try:
                import win32com.client
                speaker = win32com.client.Dispatch("SAPI.SpVoice")
                speaker.Speak(text_val)
            except Exception as ex:
                print("Speech synthesis error:", ex)
        threading.Thread(target=run_speech, daemon=True).start()

     def add_message_bubble(sender, text, is_ai=False):
        card = ctk.CTkFrame(
            chat_scroll,
            fg_color=("#F3F4F6", "#1F2937") if is_ai else ("#E0F2FE", "#053842"),
            corner_radius=12
        )
        card.pack(anchor="w" if is_ai else "e", fill="x", pady=6, padx=5)

        header = ctk.CTkLabel(
            card,
            text="🤖 SmartSearch AI" if is_ai else "🧑 You",
            font=("Segoe UI", 13, "bold"),
            text_color=("#1F2937", "#E5E7EB") if is_ai else ("#0284C7", "#7DD3FC")
        )
        header.pack(anchor="w", padx=12, pady=(8, 2))

        msg_label = ctk.CTkLabel(
            card,
            text=text,
            font=("Segoe UI", 14),
            anchor="w",
            justify="left",
            wraplength=720
        )
        msg_label.pack(anchor="w", padx=12, pady=(0, 8))

        return card, msg_label

     # Re-populate chat history bubbles when opening window
     if hasattr(self.ai, "chat_history") and self.ai.chat_history:
         for role, text in self.ai.chat_history:
             is_ai_msg = (role.lower() == "ai" or role.lower() == "assistant")
             add_message_bubble(role, text, is_ai=is_ai_msg)

     self.ai_generating = False
     self.ai_stop_flag = False

     def stop_ai():
         self.ai_stop_flag = True
         self.status.configure(text="Stopping AI generation...")

     def send():
        if getattr(self, "ai_generating", False):
            return

        question = question_box.get().strip()
        if not question:
            return

        self.ai_generating = True
        self.ai_stop_flag = False

        send_button.configure(
            text="⏸ Stop",
            fg_color="#DC2626",
            hover_color="#991B1B",
            command=stop_ai
        )

        question_box.delete(0, "end")
        add_message_bubble("You", question, is_ai=False)

        ai_card, ai_msg_label = add_message_bubble("AI", "🤖 AI is searching and thinking...", is_ai=True)

        def stream_thread():
            accumulated = ""
            scope_val = getattr(self, "chat_scope_mode", "current")
            token_gen = self.ai.stream_ask(self.ai.current_article, question, scope=scope_val)

            def safe_update(lbl, text_val):
                try:
                    if lbl.winfo_exists():
                        lbl.configure(text=text_val)
                except Exception:
                    pass

            last_update = time.time()
            for token in token_gen:
                if getattr(self, "ai_stop_flag", False):
                    accumulated += "\n\n[ ⏹ Generation Stopped by User ]"
                    win.after(0, lambda t=accumulated: safe_update(ai_msg_label, t))
                    break

                accumulated += token
                now = time.time()
                if now - last_update > 0.04:
                    last_update = now
                    curr_text = accumulated
                    win.after(0, lambda t=curr_text: safe_update(ai_msg_label, t))

            win.after(0, lambda t=accumulated: safe_update(ai_msg_label, t))

            def restore_controls():
                self.ai_generating = False
                self.ai_stop_flag = False
                try:
                    if send_button.winfo_exists():
                        send_button.configure(
                            text="Send 🚀",
                            fg_color="#053842",
                            hover_color="#5A88A4",
                            command=send
                        )
                except Exception:
                    pass

            win.after(0, restore_controls)

            def finalize():
                try:
                    if not ai_card.winfo_exists():
                        return
                    actions_bar = ctk.CTkFrame(ai_card, fg_color="transparent")
                    actions_bar.pack(anchor="e", padx=10, pady=(0, 6))

                    copy_btn = ctk.CTkButton(
                        actions_bar,
                        text="📋 Copy",
                        width=65,
                        height=22,
                        font=("Segoe UI", 11),
                        fg_color="transparent",
                        hover_color=("gray85", "gray30"),
                        text_color=("black", "white")
                    )
                    copy_btn.configure(command=lambda: copy_text(ai_msg_label.cget("text"), copy_btn))
                    copy_btn.pack(side="left", padx=3)

                    tts_btn = ctk.CTkButton(
                        actions_bar,
                        text="🔊 Read Aloud",
                        width=85,
                        height=22,
                        font=("Segoe UI", 11),
                        fg_color="transparent",
                        hover_color=("gray85", "gray30"),
                        text_color=("black", "white")
                    )
                    tts_btn.configure(command=lambda: speak_text(ai_msg_label.cget("text")))
                    tts_btn.pack(side="left", padx=3)
                except Exception:
                    pass

            win.after(0, finalize)

        threading.Thread(target=stream_thread, daemon=True).start()

     send_button = ctk.CTkButton(
        input_frame,
        text="Send 🚀",
        command=send,
        width=100,
        height=40,
        font=("Segoe UI", 14, "bold"),
        fg_color="#053842",
        hover_color="#5A88A4"
     )
     send_button.pack(side="right")
     question_box.bind("<Return>", lambda e: send())

    def exit_research_mode(self):
        topic = self.research_topic or getattr(self.ai, "current_article", None)
        self.research_mode = False
        self.research_topic = None

        if topic:
            self.search_box.delete(0, "end")
            self.search_box.insert(0, topic)
            self.search_article()
            return

        for widget in self.explorer_buttons.winfo_children():
            widget.destroy()

        self.explorer_frame.pack_forget()
        self.status.configure(
            text="Normal Search Mode: search for an article."
        )
        self.search_box.focus_set()

    def delete_current_research_pack(self):
     if not self.research_topic:
        messagebox.showinfo("Delete Research Pack", "No active research pack loaded.")
        return

     confirm = messagebox.askyesno(
        "Delete Research Pack",
        f"Are you sure you want to delete the research pack for '{self.research_topic}'?"
     )
     if confirm:
        topic_deleted = self.research_topic
        self.database.delete_research_pack(topic_deleted)
        self.exit_research_mode()
        self.status.configure(text=f"Research pack for '{topic_deleted}' deleted.")
        messagebox.showinfo("Deleted", f"Research pack for '{topic_deleted}' has been deleted from your database.")


    def build_research_pack(self):

     topic = self.ai.current_topic

     print("Topic:", topic)
     print("Type:", type(topic))

     if not topic:

      self.status.configure(
        text="Search a topic first."
    )

      return

     self.research_topic = topic

     if self.database.research_pack_exists(topic):

      self.show_research_sources(
          self.get_research_sources_for_explorer(topic)
      )

      self.research_mode = True
      self.research_topic = topic

      self.status.configure(
        text="Research Pack loaded from database."
    )

      print("Research Pack loaded locally")

      return
     if not self.check_connection():

      self.status.configure(
        text="Offline: Research Pack not available."
    )

      messagebox.showinfo(
        "Offline",
        "This research pack hasn't been downloaded yet."
    )

      return
     if not topic:
      self.status.configure(
        text="Search a topic first."
    )
      return

     self.database.delete_research_pack(topic)

     manager = SourceManager()

     downloader = ResearchDownloader()

     sources = manager.get_sources(topic)
     print("\n===== SOURCES =====")
     print(sources)
     print("===================\n")
     total = len(sources)

     count = 0

     for name, url in sources:
        print(f"\nDownloading: {name}")
        print(url)
        self.status.configure(
            text=f"Downloading {name}..."
        )

        self.window.update()

        data = downloader.download_page(
            name,
            url,
            topic
        )
        print("Returned:", data)
        if data is None:
         print(f"{name} download failed.")
         continue

        self.database.save_research_source(
            topic,
            data["name"],
            data["url"],
            data["content"],
            datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        )

        from chunker import Chunker

        chunker = Chunker()
        print("Source:", data["name"])
        print("Content length:", len(data["content"]))
        print(data["content"][:300])
        chunks = chunker.split(data["content"])

        print(f"{data['name']} -> {len(chunks)} chunks")

        for chunk in chunks:
            self.database.save_chunk(
                topic,
                data["name"],
                data["url"],
                chunk
            )

        count += 1
     self.window.configure(cursor="watch")
     self.search_button.configure(state="disabled")

     def download_worker():
          try:
              self.database.delete_research_pack(topic)
              manager = SourceManager()
              downloader = ResearchDownloader()

              sources = manager.get_sources(topic)
              count = 0

              for name, url in sources:
                  self.window.after(0, lambda n=name: self.status.configure(text=f"Downloading {n}..."))

                  data = downloader.download_page(name, url, topic)
                  if data is None:
                      continue

                  self.database.save_research_source(
                      topic,
                      data["name"],
                      data["url"],
                      data["content"],
                      datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                  )

                  from chunker import Chunker
                  chunker = Chunker()
                  chunks = chunker.split(data["content"])

                  for chunk in chunks:
                      self.database.save_chunk(topic, data["name"], data["url"], chunk)

                  count += 1

              def finalize():
                  self.window.configure(cursor="")
                  self.search_button.configure(state="normal")
                  if count == 0:
                      self.status.configure(text="Research Pack could not download any sources.")
                      return

                  self.research_mode = True
                  self.research_topic = topic
                  self.ai.current_article = topic

                  self.database.save_or_update_article(topic, f"Research Pack for {topic}")
                  full_research = self.database.get_full_research(topic)
                  self.prepare_result_box()
                  self.result_box.delete("1.0", "end")
                  self.result_box.insert("end", full_research)

                  self.status.configure(text=f"Research Pack Ready ({count} Sources)")
                  self.show_research_sources(self.get_research_sources_for_explorer(topic))

              self.window.after(0, finalize)

          except Exception as ex:
              print("Error in build_research_pack thread:", ex)
              self.window.after(0, lambda: (
                  self.window.configure(cursor=""),
                  self.search_button.configure(state="normal"),
                  self.status.configure(text="Error building research pack.")
              ))

     threading.Thread(target=download_worker, daemon=True).start()

    def get_research_sources_for_explorer(self, topic):

     sources = self.database.get_research_sources(topic)

     if sources:
        return sources

     # Research packs created before source content was persisted can still
     # be displayed by grouping their stored chunks by source name.
     grouped_sources = {}

     for chunk, source_name in self.database.get_chunks(topic):
        grouped_sources.setdefault(source_name, []).append(chunk)

     return [
        (source_name, "", "\n\n".join(chunks))
        for source_name, chunks in grouped_sources.items()
     ]

    def prepare_result_box(self):
        """Ensure result_box is in normal state so text can be updated or cleared."""
        try:
            self.result_box.configure(state="normal")
        except Exception:
            pass

    def update_reader_text(self, text):
        """Update result_box text safely while keeping it read-only for professional viewing."""
        try:
            self.result_box.configure(state="normal")
            self.result_box.delete("1.0", tk.END)
            self.result_box.insert(tk.END, text)
            self.result_box.configure(state="disabled")
        except Exception as ex:
            print("Error updating reader text:", ex)

    def show_downloaded_pdfs_library(self):
        """List all downloaded PDF files directly in the left Explorer sidebar for 100% In-App reading and AI RAG chunking."""
        try:
            for widget in self.explorer_buttons.winfo_children():
                widget.destroy()

            self.explorer_label.configure(text="📂 Saved PDF Library")

            ctk.CTkButton(
                self.explorer_buttons,
                text="📄 Upload New PDF Document",
                anchor="center",
                fg_color="#1E5631",
                hover_color="#2D8249",
                font=("Segoe UI", 12, "bold"),
                height=32,
                command=self.upload_external_pdf
            ).pack(fill="x", padx=6, pady=(8, 10))

            pdf_dir = os.path.abspath(os.path.join("data", "pdfs"))
            if not os.path.exists(pdf_dir):
                os.makedirs(pdf_dir, exist_ok=True)

            pdf_files = []
            for root_dir, dirs, files in os.walk(pdf_dir):
                for file in files:
                    if file.lower().endswith(".pdf"):
                        pdf_files.append(os.path.join(root_dir, file))

            if not pdf_files:
                ctk.CTkLabel(
                    self.explorer_buttons,
                    text="No downloaded PDFs found.\nClick 'Upload New PDF' above!",
                    font=("Segoe UI", 11, "italic"),
                    text_color="gray"
                ).pack(pady=15)
                return

            ctk.CTkLabel(
                self.explorer_buttons,
                text=f"📚 Available PDFs ({len(pdf_files)} Files):",
                font=("Segoe UI", 11, "bold")
            ).pack(anchor="w", padx=8, pady=(5, 5))

            def open_in_app_pdf(pdf_path):
                self.status.configure(text="Extracting & opening PDF in app...")
                self.window.configure(cursor="watch")

                def pdf_worker():
                    try:
                        from pdf_reader import PDFReader
                        from chunker import Chunker

                        reader = PDFReader()
                        details = reader.extract_pdf_details(pdf_path)
                        if not details or not details.get("content"):
                            self.window.after(0, lambda: (
                                self.window.configure(cursor=""),
                                messagebox.showerror("Error", "Could not extract readable text from PDF.")
                            ))
                            return

                        title = details["title"]
                        topic = f"[PDF] {title}"

                        now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                        self.database.save_or_update_article(topic, details["abstract"])
                        self.database.save_research_source(
                            topic,
                            "Saved PDF Library",
                            f"file:///{os.path.abspath(pdf_path).replace('\\', '/')}",
                            details["content"],
                            now_str
                        )

                        chunker = Chunker()
                        chunks = chunker.split(details["content"])
                        for chunk in chunks:
                            self.database.save_chunk(topic, "Saved PDF Library", f"file:///{os.path.abspath(pdf_path).replace('\\', '/')}", chunk)

                        def finalize():
                            self.research_mode = True
                            self.research_topic = topic
                            self.ai.current_article = details["content"]
                            self.ai.current_topic = topic
                            setattr(self.ai, "current_article_title", topic)

                            self.update_reader_text(details["content"])
                            self.status.configure(text=f"🟢 Opened In-App PDF: {title} ({details['page_count']} Pages)")
                            self.window.configure(cursor="")

                        self.window.after(0, finalize)
                    except Exception as ex:
                        print("Error opening PDF in app:", ex)
                        self.window.after(0, lambda: (
                            self.window.configure(cursor=""),
                            messagebox.showerror("PDF Error", f"Failed to open PDF:\n{ex}")
                        ))

                threading.Thread(target=pdf_worker, daemon=True).start()

            for pdf_path in pdf_files[:30]:
                filename = os.path.basename(pdf_path)
                display_name = filename.rsplit(".", 1)[0].replace("_", " ")[:28]
                ctk.CTkButton(
                    self.explorer_buttons,
                    text=f"📄 {display_name}",
                    anchor="w",
                    fg_color="#053842",
                    hover_color="#5A88A4",
                    height=28,
                    font=("Segoe UI", 11),
                    command=lambda p=pdf_path: open_in_app_pdf(p)
                ).pack(fill="x", padx=6, pady=2)

        except Exception as ex:
            print("Error listing PDF library:", ex)

    def upload_external_pdf(self):
        try:
            from tkinter import filedialog
            from pdf_reader import PDFReader
            from chunker import Chunker
            import shutil

            pdf_path = filedialog.askopenfilename(
                title="Select External PDF Document to Read & Extract",
                filetypes=[("PDF Documents", "*.pdf"), ("All Files", "*.*")]
            )
            if not pdf_path:
                return

            self.status.configure(text="Extracting details from uploaded PDF...")
            self.window.configure(cursor="watch")

            def upload_worker():
                try:
                    reader = PDFReader()
                    details = reader.extract_pdf_details(pdf_path)

                    if not details or not details.get("content"):
                        self.window.after(0, lambda: (
                            self.window.configure(cursor=""),
                            messagebox.showerror("Extraction Failed", "Could not extract readable text or metadata from selected PDF.")
                        ))
                        return

                    title = details["title"]
                    topic = f"[PDF] {title}"
                    clean_topic = re.sub(r'[^\w\-]', '_', topic.strip())

                    target_dir = os.path.abspath(os.path.join("data", "pdfs", clean_topic))
                    os.makedirs(target_dir, exist_ok=True)
                    local_pdf_copy = os.path.join(target_dir, os.path.basename(pdf_path))
                    if not os.path.exists(local_pdf_copy):
                        shutil.copy2(pdf_path, local_pdf_copy)

                    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                    self.database.save_or_update_article(topic, details["abstract"])
                    self.database.save_research_source(
                        topic,
                        "Uploaded PDF Document",
                        f"file:///{os.path.abspath(pdf_path).replace('\\', '/')}",
                        details["content"],
                        now_str
                    )

                    chunker = Chunker()
                    chunks = chunker.split(details["content"])
                    for chunk in chunks:
                        self.database.save_chunk(topic, "Uploaded PDF Document", f"file:///{os.path.abspath(pdf_path).replace('\\', '/')}", chunk)

                    def finalize():
                        self.research_mode = True
                        self.research_topic = topic
                        self.ai.current_article = details["content"]
                        self.ai.current_topic = topic

                        self.search_box.delete(0, tk.END)
                        self.search_box.insert(0, topic)

                        self.update_reader_text(details["content"])
                        self.status.configure(text=f"🟢 Extracted PDF: {title} ({details['page_count']} Pages, {details['file_size']})")
                        self.window.configure(cursor="")
                        messagebox.showinfo("PDF Extracted Successfully!", f"Document: {title}\nPages: {details['page_count']}\nSize: {details['file_size']}\n\nYou can now read it, ask AI Chat questions, and generate Study Notes!")

                    self.window.after(0, finalize)
                except Exception as ex:
                    print("Error processing upload PDF:", ex)
                    self.window.after(0, lambda: (
                        self.window.configure(cursor=""),
                        messagebox.showerror("Upload Error", f"Failed to process PDF:\n{ex}")
                    ))

            threading.Thread(target=upload_worker, daemon=True).start()
        except Exception as ex:
            print("Error uploading PDF:", ex)



    def copy_citation(self, title, authors="", year="2024", journal="Scientific Repository", url="", style="IEEE"):
        try:
            from citation_formatter import CitationFormatter
            formatter = CitationFormatter()
            if style == "APA":
                formatted = formatter.format_apa(title, authors, year, journal, url)
            elif style == "BibTeX":
                formatted = formatter.format_bibtex(title, authors, year, journal, url)
            else:
                formatted = formatter.format_ieee(title, authors, year, journal, url)

            self.window.clipboard_clear()
            self.window.clipboard_append(formatted)
            self.status.configure(text=f"🟢 Copied {style} Citation to Clipboard!")
            messagebox.showinfo(f"{style} Citation Copied", f"Copied {style} Citation to Clipboard:\n\n{formatted}")
        except Exception as ex:
            print("Error copying citation:", ex)

    def show_citation_dialog(self):
        """Open a dedicated citation modal for the currently opened document offering IEEE, APA, BibTeX, and Copy All options."""
        topic = getattr(self.ai, "current_article_title", None) or getattr(self, "research_topic", None) or getattr(self.ai, "current_article", None)
        if not topic or isinstance(topic, list):
            topic = getattr(self.ai, "current_topic", None) or self.search_box.get().strip()

        if not topic or not str(topic).strip() or str(topic).strip().startswith("Research Pack for"):
            topic = self.search_box.get().strip()

        if not topic or len(str(topic).strip()) == 0:
            messagebox.showinfo("Cite Document", "Please open or search an article/document first to generate citations.")
            return

        doc_title = str(topic).strip()
        url = ""
        authors = "Academic Researchers"
        year = "2024"
        journal = "SmartSearch Library"

        sources = self.database.get_research_sources(doc_title)
        if sources and len(sources) > 0:
            if len(sources[0]) > 1 and sources[0][1]:
                url = sources[0][1]
            if len(sources[0]) > 0 and sources[0][0]:
                journal = sources[0][0]

        try:
            from citation_formatter import CitationFormatter
            formatter = CitationFormatter()

            ieee_text = formatter.format_ieee(doc_title, authors, year, journal, url)
            apa_text = formatter.format_apa(doc_title, authors, year, journal, url)
            bibtex_text = formatter.format_bibtex(doc_title, authors, year, journal, url)

            dialog = ctk.CTkToplevel(self.window)
            dialog.title(f"🎓 Academic Citations — {doc_title[:30]}")
            dialog.geometry("640x540")
            dialog.grab_set()
            dialog.focus_force()

            ctk.CTkLabel(
                dialog,
                text=f"Academic Citations for Opened Document:\n\"{doc_title}\"",
                font=("Segoe UI", 13, "bold"),
                wraplength=600,
                justify="center"
            ).pack(pady=(15, 10))

            # IEEE Box
            ieee_frame = ctk.CTkFrame(dialog)
            ieee_frame.pack(fill="x", padx=20, pady=5)
            ctk.CTkLabel(ieee_frame, text="IEEE Format:", font=("Segoe UI", 11, "bold")).pack(anchor="w", padx=10, pady=(5, 0))
            ieee_entry = ctk.CTkEntry(ieee_frame, width=460)
            ieee_entry.insert(0, ieee_text)
            ieee_entry.pack(side="left", padx=10, pady=5, expand=True, fill="x")
            ctk.CTkButton(
                ieee_frame,
                text="📋 Copy IEEE",
                width=100,
                fg_color="#005B96",
                hover_color="#003B63",
                command=lambda: (
                    self.window.clipboard_clear(),
                    self.window.clipboard_append(ieee_text),
                    self.status.configure(text="🟢 IEEE Citation Copied!"),
                    messagebox.showinfo("Copied", "IEEE Citation copied to clipboard!")
                )
            ).pack(side="right", padx=10, pady=5)

            # APA Box
            apa_frame = ctk.CTkFrame(dialog)
            apa_frame.pack(fill="x", padx=20, pady=5)
            ctk.CTkLabel(apa_frame, text="APA 7th Edition Format:", font=("Segoe UI", 11, "bold")).pack(anchor="w", padx=10, pady=(5, 0))
            apa_entry = ctk.CTkEntry(apa_frame, width=460)
            apa_entry.insert(0, apa_text)
            apa_entry.pack(side="left", padx=10, pady=5, expand=True, fill="x")
            ctk.CTkButton(
                apa_frame,
                text="📋 Copy APA",
                width=100,
                fg_color="#053842",
                hover_color="#5A88A4",
                command=lambda: (
                    self.window.clipboard_clear(),
                    self.window.clipboard_append(apa_text),
                    self.status.configure(text="🟢 APA Citation Copied!"),
                    messagebox.showinfo("Copied", "APA Citation copied to clipboard!")
                )
            ).pack(side="right", padx=10, pady=5)

            # BibTeX Box
            bib_frame = ctk.CTkFrame(dialog)
            bib_frame.pack(fill="x", padx=20, pady=5)
            ctk.CTkLabel(bib_frame, text="BibTeX Format (LaTeX):", font=("Segoe UI", 11, "bold")).pack(anchor="w", padx=10, pady=(5, 0))
            bib_textbox = ctk.CTkTextbox(bib_frame, height=75)
            bib_textbox.insert("1.0", bibtex_text)
            bib_textbox.pack(side="left", padx=10, pady=5, expand=True, fill="both")
            ctk.CTkButton(
                bib_frame,
                text="📋 Copy BibTeX",
                width=100,
                fg_color="#1E5631",
                hover_color="#2D8249",
                command=lambda: (
                    self.window.clipboard_clear(),
                    self.window.clipboard_append(bibtex_text),
                    self.status.configure(text="🟢 BibTeX Entry Copied!"),
                    messagebox.showinfo("Copied", "BibTeX Entry copied to clipboard!")
                )
            ).pack(side="right", padx=10, pady=5)

            # Bottom Action Bar
            bottom_frame = ctk.CTkFrame(dialog, fg_color="transparent")
            bottom_frame.pack(fill="x", padx=20, pady=15)

            def copy_all():
                all_text = f"=== IEEE ===\n{ieee_text}\n\n=== APA 7th Edition ===\n{apa_text}\n\n=== BibTeX ===\n{bibtex_text}"
                self.window.clipboard_clear()
                self.window.clipboard_append(all_text)
                self.status.configure(text="🟢 All Citations Copied!")
                messagebox.showinfo("Copied All", "All citations (IEEE, APA, BibTeX) copied to clipboard!")

            ctk.CTkButton(
                bottom_frame,
                text="📋 Copy All Citations",
                fg_color="#005B96",
                hover_color="#003B63",
                font=("Segoe UI", 11, "bold"),
                command=copy_all
            ).pack(side="left", padx=10)

            ctk.CTkButton(
                bottom_frame,
                text="Close",
                fg_color="gray40",
                command=dialog.destroy
            ).pack(side="right", padx=10)
        except Exception as ex:
            print("Error showing citation modal:", ex)
            messagebox.showerror("Error", f"Failed to generate citations:\n{ex}")

    def show_research_sources(self, sources):

     self.explorer_frame.pack(
        side="left",
        fill="y",
        padx=(0, 5),
        before=self.results_container
     )

     for widget in self.explorer_buttons.winfo_children():
        widget.destroy()

     self.explorer_label.configure(text="📖 Research Sources")

     for source_name, source_url, content in sources:
        if source_name not in ["arXiv Papers", "OpenAlex Literature", "PubMed Central", "Semantic Scholar"]:
            continue

        sections = self.explorer.extract_sections(content)
        valid_sections = [s for s in sections if s.get("title", "").strip() and s.get("title", "").strip() not in ["Full Article", "="] and len(s.get("title", "").strip()) >= 2]
        paper_count = len(valid_sections)

        header_label = f"📄 {source_name} ({paper_count} Papers)" if paper_count > 0 else f"📄 {source_name}"

        ctk.CTkButton(
            self.explorer_buttons,
            text=header_label,
            anchor="w",
            fg_color="#053842",
            hover_color="#5A88A4",
            command=lambda n=source_name, u=source_url, c=content: (
                self.open_research_source(n, u, c)
            )
        ).pack(fill="x", pady=(6, 2))

        for section in valid_sections:
            sec_title = section.get("title", "").strip()
            ctk.CTkButton(
                self.explorer_buttons,
                text=f"   ↳ {sec_title[:30]}",
                anchor="w",
                fg_color="transparent",
                hover_color=("gray85", "gray25"),
                text_color=("black", "white"),
                command=lambda s=section: self.open_section(s)
            ).pack(fill="x", padx=8, pady=1)

    def open_research_source(self, source_name, source_url, content):

     self.prepare_result_box()
     self.result_box.delete("1.0", "end")

     self.result_box.insert("end", f"=== SOURCE: {source_name} ===\n")

     if source_url:
         self.result_box.insert("end", f"URL: {source_url}\n")

     self.result_box.insert("end", "\n")

     self.result_box.insert("end", content)

    def show_research_explorer(self, sections):

     for widget in self.explorer_buttons.winfo_children():
        widget.destroy()

     self.explorer_frame.pack(
        side="left",
        fill="y",
        padx=10,
        pady=10
    )

     for section in sections:

        btn = ctk.CTkButton(
            self.explorer_buttons,
            text=section["title"],
            anchor="w",
            fg_color="#053842",
            hover_color="#5A88A4",
            command=lambda s=section: self.open_section(s)
        )

        btn.pack(
            fill="x",
            pady=2
        )

    def open_section(self, section):

     self.prepare_result_box()
     self.result_box.delete(
        "1.0",
        "end"
    )

     self.result_box.insert(
        "end",
        section["title"] + "\n\n"
    )

     self.result_box.insert(
        "end",
        section["content"]
    )

    def create_popup_menu(self):

     self.popup_menu = tk.Menu(
        self.window,
        tearoff=0
    )

     self.popup_menu.add_command(
        label="📜 History",
        command=self.show_history
    )

     self.popup_menu.add_command(
        label="⭐ Favorites",
        command=self.show_favorites
    )

     self.popup_menu.add_command(
        label="🔥 Most Viewed",
        command=self.show_most_viewed
    )

     self.popup_menu.add_separator()

     self.popup_menu.add_command(
        label="📊 Statistics",
        command=self.show_statistics
    )

     self.popup_menu.add_command(
        label="🌙 Dark Mode",
        command=self.toggle_dark_mode
    )

     self.popup_menu.add_separator()

     self.popup_menu.add_command(
        label="⭐ Add to Favorites",
        command=self.add_to_favorites
    )

     self.popup_menu.add_separator()

     self.popup_menu.add_command(
        label="📝 Export MS Word",
        command=self.export_docx
    )

     self.popup_menu.add_command(
        label="📄 Export PDF",
        command=self.export_pdf
    )

     self.popup_menu.add_command(
        label="📜 Export TXT",
        command=self.export_txt
    )

     self.popup_menu.add_separator()

     self.popup_menu.add_command(
        label="🎓 Academic Citation...",
        command=self.show_citation_dialog
    )

    def show_menu(self):

     x = self.menu_button.winfo_rootx()
     y = self.menu_button.winfo_rooty() + self.menu_button.winfo_height()

     self.popup_menu.post(x, y)


